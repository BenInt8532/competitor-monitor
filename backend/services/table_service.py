"""
Разбор сравнительных таблиц (XLSX, CSV, таблицы внутри Word).

Файл может быть не «одним документом про обменник», а списком конкурентов:
в колонках — название, ссылка, курс, способы выплаты, заметки.
Приложение само решает, что перед ним, и раскладывает строки по профилям.

Правило: если нашлась хотя бы одна ссылка (или ячейка-домен вроде grumbot.net) —
это таблица конкурентов. Каждая строка = один обменник.
Данные из ячеек пишем в память сразу; ссылки потом разбираются как обычные URL.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from docx import Document

from backend.utils import DOMAIN_LIKE, competitor_name_from_url, domain_from_url, normalize_url

NAME_HINTS = ("название", "имя", "name", "обменник", "competitor", "сервис", "exchanger", "бренд")
URL_HINTS = ("url", "сайт", "site", "ссылка", "домен", "domain", "link", "website", "лендинг")
RATE_HINTS = ("курс", "rate", "usdt", "цена", "котир")
PAYOUT_HINTS = ("выплат", "payout", "способ", "реквизит", "наличн")

URL_IN_TEXT = re.compile(
    r"(https?://[^\s,;|]+|(?:www\.)?[a-z0-9][a-z0-9\-.]+\.[a-z]{2,6}(?:/[^\s,;|]*)?)",
    re.IGNORECASE,
)


@dataclass
class TableRow:
    competitor_name: str
    url: Optional[str] = None
    exchanger_rate: Optional[float] = None
    payout_methods: List[str] = field(default_factory=list)
    notes: str = ""
    extras: Dict[str, str] = field(default_factory=dict)
    row_index: int = 0
    url_source: Optional[str] = None  # "файл" или "поиск" — откуда взялась ссылка


@dataclass
class TableParse:
    rows: List[TableRow]
    headers: List[str]
    is_competitor_table: bool
    raw_preview: str = ""


class TableService:
    def parse(self, content: bytes, filename: str, kind: str) -> TableParse:
        suffix = Path(filename or "").suffix.lower()
        if kind == "xlsx" or suffix in {".xlsx", ".xlsm"}:
            matrix = self._from_xlsx(content)
        elif kind == "csv" or suffix == ".csv":
            matrix = self._from_csv(content)
        elif kind == "docx" or suffix == ".docx":
            matrix = self._from_docx_tables(content)
        else:
            matrix = []

        return self._interpret(matrix)

    # --- чтение в матрицу ячеек ---

    @staticmethod
    def _from_xlsx(content: bytes) -> List[List[str]]:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        matrix: List[List[str]] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    matrix.append(cells)
        workbook.close()
        return matrix

    @staticmethod
    def _from_csv(content: bytes) -> List[List[str]]:
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = content.decode("utf-8", errors="ignore")

        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
            if sample.count(";") > sample.count(","):
                dialect.delimiter = ";"

        return [
            [cell.strip() for cell in row]
            for row in csv.reader(io.StringIO(text), dialect)
            if any(cell.strip() for cell in row)
        ]

    @staticmethod
    def _from_docx_tables(content: bytes) -> List[List[str]]:
        document = Document(io.BytesIO(content))
        matrix: List[List[str]] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    matrix.append(cells)
        return matrix

    # --- понимание, что в таблице ---

    def _interpret(self, matrix: List[List[str]]) -> TableParse:
        if not matrix:
            return TableParse(rows=[], headers=[], is_competitor_table=False)

        headers, data = self._split_header(matrix)
        roles = self._classify_columns(headers)

        rows: List[TableRow] = []
        for index, cells in enumerate(data, start=1):
            parsed = self._row_from_cells(cells, headers, roles, index)
            if parsed:
                rows.append(parsed)

        urls = sum(1 for row in rows if row.url)
        named = sum(1 for row in rows if row.competitor_name)
        is_table = urls >= 1 or named >= 2

        preview_lines = [" | ".join(headers)] if headers else []
        preview_lines += [" | ".join(filter(None, [r.competitor_name, r.url or "", r.notes])) for r in rows[:12]]

        return TableParse(
            rows=rows,
            headers=headers,
            is_competitor_table=is_table,
            raw_preview="\n".join(preview_lines),
        )

    def _split_header(self, matrix: List[List[str]]) -> tuple:
        first = matrix[0]
        joined = " ".join(first).lower()
        looks_like_header = any(
            hint in joined
            for hint in NAME_HINTS + URL_HINTS + RATE_HINTS + PAYOUT_HINTS
        )
        if looks_like_header:
            return [cell or f"колонка {i + 1}" for i, cell in enumerate(first)], matrix[1:]
        # заголовка нет — колонки безымянные, всё равно ищем ссылки по ячейкам
        width = max(len(row) for row in matrix)
        return [f"колонка {i + 1}" for i in range(width)], matrix

    @staticmethod
    def _classify_columns(headers: List[str]) -> Dict[int, str]:
        roles: Dict[int, str] = {}
        used = set()
        for index, header in enumerate(headers):
            low = header.lower()
            if any(h in low for h in URL_HINTS):
                roles[index] = "url"
                used.add("url")
            elif any(h in low for h in NAME_HINTS):
                roles[index] = "name"
                used.add("name")
            elif any(h in low for h in RATE_HINTS):
                roles[index] = "rate"
                used.add("rate")
            elif any(h in low for h in PAYOUT_HINTS):
                roles[index] = "payout"
                used.add("payout")
        return roles

    def _row_from_cells(
        self,
        cells: List[str],
        headers: List[str],
        roles: Dict[int, str],
        row_index: int,
    ) -> Optional[TableRow]:
        url = None
        name = ""
        rate = None
        payouts: List[str] = []
        extras: Dict[str, str] = {}

        for index, raw in enumerate(cells):
            value = (raw or "").strip()
            if not value:
                continue
            role = roles.get(index)
            header = headers[index] if index < len(headers) else f"колонка {index + 1}"

            found_url = self._extract_url(value)
            if role == "url" or (not url and found_url):
                url = url or found_url
                if found_url and value != found_url:
                    extras[header] = value
                continue
            if role == "name":
                name = value
                continue
            if role == "rate":
                rate = self._parse_rate(value)
                extras[header] = value
                continue
            if role == "payout":
                payouts.extend(self._split_list(value))
                extras[header] = value
                continue

            maybe_rate = self._parse_rate(value) if self._looks_like_rate(value) else None
            if maybe_rate is not None and rate is None:
                rate = maybe_rate
                extras[header] = value
                continue

            extras[header] = value
            if not name and not found_url and not DOMAIN_LIKE.match(value):
                # первое текстовое поле, не ссылка — кандидат в название
                if len(value) <= 60:
                    name = name or value

        if not name and url:
            name = competitor_name_from_url(url)
        if not name and extras:
            name = next(iter(extras.values()), "")[:60]

        if not name and not url:
            return None

        notes = " · ".join(f"{key}: {val}" for key, val in extras.items() if val)
        return TableRow(
            competitor_name=name[:80],
            url=url,
            exchanger_rate=rate,
            payout_methods=payouts,
            notes=notes,
            extras=extras,
            row_index=row_index,
            url_source="файл" if url else None,
        )

    def parse_text(self, text: str) -> TableParse:
        """
        Свободный текст: одно описание, список ссылок или одни названия.
        Приложение само решает, что перед ним.
        """
        raw_lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
        if not raw_lines:
            return TableParse(rows=[], headers=[], is_competitor_table=False, raw_preview="")

        # одна строка «GrumBot, CoinDrop, StoreBucks» — это список имён, не предложение
        if len(raw_lines) == 1:
            parts = [part.strip() for part in re.split(r"[,;]", raw_lines[0]) if part.strip()]
            if len(parts) >= 2 and all(self._looks_like_name_line(part) for part in parts):
                raw_lines = parts

        rows: List[TableRow] = []
        prose = 0
        for index, line in enumerate(raw_lines, start=1):
            cleaned = self._strip_bullet(line)
            row = self._row_from_free_line(cleaned, index)
            if row:
                rows.append(row)
            else:
                prose += 1

        # два и больше пунктов — список. Один короткий пункт без прозы — тоже
        # (одна ссылка или одно имя). Длинное описание с одной ссылкой внутри — не список.
        total_chars = sum(len(line) for line in raw_lines)
        is_list = len(rows) >= 2 or (
            len(rows) == 1 and prose == 0 and total_chars < 220
        )

        preview = "\n".join(
            " | ".join(filter(None, [r.competitor_name, r.url or "", r.notes]))
            for r in rows[:12]
        )
        return TableParse(
            rows=rows if is_list else [],
            headers=[],
            is_competitor_table=is_list,
            raw_preview=preview,
        )

    def first_url_in_text(self, text: str) -> Optional[str]:
        return self._extract_url(text or "")

    def _row_from_free_line(self, line: str, index: int) -> Optional[TableRow]:
        url = self._extract_url(line)
        rest = line
        if url:
            rest = URL_IN_TEXT.sub(" ", line).strip(" -—–|,;")
        rate = self._parse_rate(line)

        name = ""
        notes = rest
        if rest:
            split = re.split(r"\s+[—–-]\s+|:\s+", rest, maxsplit=1)
            if self._looks_like_name_line(split[0]):
                name = split[0].strip()
                notes = split[1].strip() if len(split) > 1 else ""
            elif url and len(rest) <= 80:
                name = rest
                notes = ""

        if not name and url:
            name = competitor_name_from_url(url)
        if not name and self._looks_like_name_line(line):
            name = self._strip_bullet(line)
            notes = ""
        if not name:
            return None

        return TableRow(
            competitor_name=name[:80],
            url=url,
            exchanger_rate=rate,
            notes=notes,
            row_index=index,
            url_source="файл" if url else None,
        )

    @staticmethod
    def _strip_bullet(line: str) -> str:
        return re.sub(r"^(\d+[.)]\s+|[-*•—]\s+)", "", line).strip()

    def _looks_like_name_line(self, value: str) -> bool:
        text = self._strip_bullet(value or "")
        if not text or len(text) > 50:
            return False
        if self._extract_url(text) and URL_IN_TEXT.sub("", text).strip() == "":
            return True
        if DOMAIN_LIKE.match(text):
            return True
        words = text.split()
        if not 1 <= len(words) <= 5:
            return False
        if text.endswith(".") and len(words) > 2:
            return False
        low = text.lower()
        if low.startswith(("это ", "обменник ", "сервис ", "компания ", "если ", "при ", "мы ")):
            return False
        return bool(re.search(r"[a-zA-Zа-яА-Я]", text))

    @staticmethod
    def _extract_url(value: str) -> Optional[str]:
        match = URL_IN_TEXT.search(value or "")
        if not match:
            return None
        candidate = normalize_url(match.group(1).rstrip(").,]"))
        return candidate if domain_from_url(candidate) else None

    @staticmethod
    def _parse_rate(value: str) -> Optional[float]:
        cleaned = (value or "").replace("\xa0", " ").replace(",", ".")
        match = re.search(r"(\d+(?:\.\d+)?)", cleaned)
        if not match:
            return None
        number = float(match.group(1))
        # курс USDT/RUB живёт примерно в этом диапазоне; 1000 — это уже сумма
        if 20 <= number <= 200:
            return number
        return None

    @staticmethod
    def _looks_like_rate(value: str) -> bool:
        low = value.lower()
        return any(token in low for token in ("курс", "usdt", "₽", "руб")) or bool(
            re.fullmatch(r"\d+[.,]\d{1,4}", value.strip())
        )

    @staticmethod
    def _split_list(value: str) -> List[str]:
        parts = re.split(r"[,;/|]| и ", value)
        return [part.strip() for part in parts if part.strip()]


table_service = TableService()
